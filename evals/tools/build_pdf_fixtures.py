"""Build the controlled Sprint 2 PDF fixtures through LibreOffice.

The DOCX intermediate gives the PDF both correct Persian shaping and logical
Unicode text that the production pypdf extractor can recover.
"""

from __future__ import annotations

import subprocess
import tempfile
from pathlib import Path
from xml.etree import ElementTree

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from pypdf import PdfReader, PdfWriter


ROOT = Path(__file__).resolve().parents[1]
SOFFICE = Path(
    "/Users/a.abnar/.cache/codex-runtimes/codex-primary-runtime/"
    "dependencies/bin/override/soffice"
)
FIXTURES = (
    ("source/sprint-2-fa-operations.html", "sprint-2-fa-operations.pdf"),
    ("source/sprint-2-en-product.html", "sprint-2-en-product.pdf"),
)


def main() -> None:
    if not SOFFICE.exists():
        raise SystemExit(f"LibreOffice was not found at {SOFFICE}")
    fixture_root = ROOT / "fixtures"
    with tempfile.TemporaryDirectory(prefix="docqa-pdf-fixtures-") as temp_name:
        temp_dir = Path(temp_name)
        for source_name, output_name in FIXTURES:
            source = (fixture_root / source_name).resolve()
            docx_path = temp_dir / f"{Path(source_name).stem}.docx"
            build_docx(source, docx_path)
            subprocess.run(
                (
                    str(SOFFICE),
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(fixture_root),
                    str(docx_path),
                ),
                check=True,
            )
            generated = fixture_root / docx_path.with_suffix(".pdf").name
            output = fixture_root / output_name
            if generated != output:
                generated.replace(output)
            canonicalize_pdf(output, temp_dir / f"canonical-{output_name}")
            print(output)


def build_docx(source: Path, output: Path) -> None:
    markup = source.read_text(encoding="utf-8")
    root = ElementTree.fromstring(markup[markup.index("<html") :])
    body = root.find("body")
    if body is None:
        raise ValueError(f"Missing body in {source}")
    language = root.attrib.get("lang", "en")
    rtl = language == "fa"
    document = Document()
    section = document.sections[0]
    section.page_height = Mm(297)
    section.page_width = Mm(210)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    normal = document.styles["Normal"]
    normal.font.name = "Tahoma" if rtl else "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.25

    pages = [child for child in body if child.tag == "section"]
    for page_index, page in enumerate(pages):
        for element in page:
            if element.tag == "header":
                for child in element:
                    add_element(document, child, rtl=rtl)
            else:
                add_element(document, element, rtl=rtl)
        if page_index < len(pages) - 1:
            document.add_page_break()
    document.save(output)


def add_element(document: Document, element: ElementTree.Element, *, rtl: bool) -> None:
    if element.tag == "h1":
        paragraph = document.add_heading(level=1)
    elif element.tag == "h2":
        paragraph = document.add_heading(level=2)
    elif element.tag in {"p", "div", "footer"}:
        paragraph = document.add_paragraph()
    else:
        return
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    if rtl:
        properties = paragraph._p.get_or_add_pPr()
        properties.append(OxmlElement("w:bidi"))
    if element.tag == "footer" or element.attrib.get("class") == "meta":
        paragraph.style = document.styles["Caption"]
    add_text_runs(paragraph, element, rtl=rtl)


def add_text_runs(paragraph: object, element: ElementTree.Element, *, rtl: bool) -> None:
    if element.text:
        configure_run(paragraph.add_run(element.text), rtl=rtl)
    for child in element:
        child_rtl = rtl and child.attrib.get("dir") != "ltr"
        if child.text:
            configure_run(paragraph.add_run(child.text), rtl=child_rtl)
        if child.tail:
            configure_run(paragraph.add_run(child.tail), rtl=rtl)


def configure_run(run: object, *, rtl: bool) -> None:
    run.font.name = "Tahoma" if rtl else "Arial"
    if rtl:
        properties = run._r.get_or_add_rPr()
        rtl_element = OxmlElement("w:rtl")
        rtl_element.set(qn("w:val"), "1")
        properties.append(rtl_element)


def canonicalize_pdf(path: Path, temporary_output: Path) -> None:
    reader = PdfReader(path)
    writer = PdfWriter()
    writer.clone_document_from_reader(reader)
    writer.metadata = None
    writer.add_metadata(
        {
            "/Title": path.stem,
            "/Author": "Document Q&A evaluation",
            "/Creator": "versioned fixture builder",
            "/Producer": "pypdf",
            "/CreationDate": "D:20260926000000Z",
            "/ModDate": "D:20260926000000Z",
        }
    )
    writer.generate_file_identifiers()
    with temporary_output.open("wb") as target:
        writer.write(target)
    temporary_output.replace(path)


if __name__ == "__main__":
    main()
